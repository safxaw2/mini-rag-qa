import os
import sys

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import tempfile

import docx

from src.parsers.base import ParseResult
from src.parsers.chunker import TextChunker
from src.parsers.cleaner import TextCleaner
from src.parsers.table_parser import TableProcessor
from src.parsers.word_paser import Wordarser


# ---------------------------------------------------------------------------
# 测试辅助：用 python-docx 现场生成文档，避免依赖 data/ 目录中的样例文件
# ---------------------------------------------------------------------------

def _make_doc(paragraphs=None, tables=None):
    """创建一个内存中的 docx 文档"""
    doc = docx.Document()
    for text in paragraphs or []:
        doc.add_paragraph(text)
    for rows in tables or []:
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for row_index, row in enumerate(rows):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
    return doc


def _save_temp(doc):
    """把文档保存到临时文件，返回路径（调用方负责删除）"""
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)
    return path


TABLE_ROWS = [["名称", "数量"], ["苹果", "3"], ["香蕉", "5"]]


class TestSupportedExtensions:
    """supported_extensions 方法"""

    def test_returns_docx_extensions(self):
        assert Wordarser().supported_extensions() == [".docx", ".DOCX"]


class TestInit:
    """Wordarser 初始化：组件可注入，缺省时使用默认实现"""

    def test_default_components(self):
        parser = Wordarser()
        assert isinstance(parser.chunker, TextChunker)
        assert isinstance(parser.cleaner, TextCleaner)
        assert isinstance(parser.table_processor, TableProcessor)

    def test_injected_components_are_kept(self):
        chunker = TextChunker(max_chunk_size=10)
        cleaner = TextCleaner()
        table_processor = TableProcessor()
        parser = Wordarser(
            chunker=chunker, cleaner=cleaner, table_processor=table_processor
        )
        assert parser.chunker is chunker
        assert parser.cleaner is cleaner
        assert parser.table_processor is table_processor


class TestExtractTables:
    """_extract_tables 方法：docx 表格 -> 结构化字典列表"""

    def test_no_tables_returns_empty_list(self):
        doc = _make_doc(["只有正文"])
        assert Wordarser()._extract_tables(doc) == []

    def test_table_is_converted_to_markdown(self):
        doc = _make_doc(tables=[TABLE_ROWS])
        tables = Wordarser()._extract_tables(doc)

        assert len(tables) == 1
        table = tables[0]
        assert table["page"] == 1
        assert table["row"] == 3
        assert table["cols"] == 2
        assert table["markdown"] == "\n".join(
            [
                "|名称|数量|",
                "|---|---|",
                "|苹果|3|",
                "|香蕉|5|",
            ]
        )

    def test_cells_are_stripped(self):
        doc = _make_doc(tables=[[[" a ", " b "], [" c ", " d "]]])
        markdown = Wordarser()._extract_tables(doc)[0]["markdown"]
        assert markdown.splitlines()[0] == "|a|b|"
        assert markdown.splitlines()[2] == "|c|d|"

    def test_multiple_tables(self):
        doc = _make_doc(tables=[[["A"], ["1"]], [["B"], ["2"]]])
        tables = Wordarser()._extract_tables(doc)
        assert len(tables) == 2
        assert [t["row"] for t in tables] == [2, 2]
        assert [t["cols"] for t in tables] == [1, 1]


class TestParse:
    """parse 方法：按路径解析 docx 文件"""

    def test_returns_parse_result(self):
        path = _save_temp(_make_doc(["第一段", "第二段"]))
        try:
            result = Wordarser().parse(path)
            assert isinstance(result, ParseResult)
        finally:
            os.remove(path)

    def test_file_name_is_basename(self):
        path = _save_temp(_make_doc(["正文"]))
        try:
            result = Wordarser().parse(path)
            assert result.file_name == os.path.basename(path)
        finally:
            os.remove(path)

    def test_raw_text_joins_non_empty_paragraphs_with_blank_line(self):
        path = _save_temp(_make_doc(["第一段", "   ", "第二段"]))
        try:
            result = Wordarser().parse(path)
            assert result.raw_text == "第一段\n\n第二段"
        finally:
            os.remove(path)

    def test_total_pages_counts_paragraphs(self):
        path = _save_temp(_make_doc(["一", "二", "三"]))
        try:
            result = Wordarser().parse(path)
            assert result.total_pages == 3
        finally:
            os.remove(path)

    def test_metadata_marks_python_docx_parser(self):
        path = _save_temp(_make_doc(["正文"]))
        try:
            result = Wordarser().parse(path)
            assert result.metadata == {"parser": "python-docx"}
        finally:
            os.remove(path)

    def test_cleaned_text_drops_page_number_lines(self):
        # 清洗器会删除“第 3 页”这类无意义行，原始文本仍保留
        path = _save_temp(_make_doc(["正文内容", "第 3 页", "更多内容"]))
        try:
            parser = Wordarser(chunker=TextChunker(min_chunk_size=1))
            result = parser.parse(path)
            assert "第 3 页" in result.raw_text
            contents = [chunk["content"] for chunk in result.chunks]
            assert not any("第 3 页" in content for content in contents)
        finally:
            os.remove(path)

    def test_empty_document_produces_no_chunks(self):
        path = _save_temp(_make_doc())
        try:
            result = Wordarser().parse(path)
            assert result.raw_text == ""
            assert result.total_pages == 0
            assert result.table_count == 0
            assert result.chunks == []
            assert result.tables == []
        finally:
            os.remove(path)

    def test_table_count_and_tables(self):
        path = _save_temp(_make_doc(["正文"], tables=[TABLE_ROWS]))
        try:
            result = Wordarser().parse(path)
            assert result.table_count == 1
            assert len(result.tables) == 1
            assert "|名称|数量|" in result.tables[0]["markdown"]
        finally:
            os.remove(path)

    def test_tables_are_attached_as_chunks(self):
        path = _save_temp(_make_doc(["一段足够长的正文内容。"], tables=[TABLE_ROWS]))
        try:
            parser = Wordarser(chunker=TextChunker(min_chunk_size=1))
            result = parser.parse(path)
            table_chunks = [c for c in result.chunks if c["type"] == "table"]
            # 表格的每一行数据会被 merge_tables 合并后单独成块
            assert len(table_chunks) == 2
            assert [c["index"] for c in result.chunks] == list(
                range(len(result.chunks))
            )
        finally:
            os.remove(path)

    def test_custom_cleaner_is_used(self):
        class DummyCleaner(TextCleaner):
            def clean(self, text):
                return "CLEANED"

        path = _save_temp(_make_doc(["原文"]))
        try:
            parser = Wordarser(
                cleaner=DummyCleaner(), chunker=TextChunker(min_chunk_size=1)
            )
            result = parser.parse(path)
            assert result.chunks[0]["content"] == "CLEANED"
        finally:
            os.remove(path)


class TestParseFromBytes:
    """parse_from_bytes 方法：从字节流解析 docx"""

    def test_parses_bytes(self):
        path = _save_temp(_make_doc(["第一段", "第二段"]))
        try:
            with open(path, "rb") as f:
                file_bytes = f.read()
            result = Wordarser().parse_from_bytes(file_bytes, "sample.docx")
            assert result.file_name == "sample.docx"
            assert result.raw_text == "第一段\n\n第二段"
            assert result.total_pages == 2
        finally:
            os.remove(path)

    def test_matches_parse_result(self):
        path = _save_temp(_make_doc(["正文"], tables=[TABLE_ROWS]))
        try:
            with open(path, "rb") as f:
                file_bytes = f.read()
            parser = Wordarser(chunker=TextChunker(min_chunk_size=1))
            from_path = parser.parse(path)
            from_bytes = parser.parse_from_bytes(file_bytes, "sample.docx")

            assert from_bytes.chunks == from_path.chunks
            assert from_bytes.tables == from_path.tables
            assert from_bytes.table_count == from_path.table_count
            assert from_bytes.raw_text == from_path.raw_text
        finally:
            os.remove(path)

    def test_metadata_includes_file_name(self):
        path = _save_temp(_make_doc(["正文"]))
        try:
            with open(path, "rb") as f:
                file_bytes = f.read()
            result = Wordarser().parse_from_bytes(file_bytes, "报告.docx")
            assert result.metadata == {
                "parser": "python-docx",
                "file_name": "报告.docx",
            }
        finally:
            os.remove(path)

    def test_invalid_bytes_raises(self):
        # 非 docx 内容应当由 python-docx 抛出异常
        try:
            Wordarser().parse_from_bytes(b"not a docx file", "bad.docx")
        except Exception:
            pass
        else:
            raise AssertionError("非法字节流应当抛出异常")


if __name__ == "__main__":
    # 无 pytest 时也可直接运行：手动收集并执行所有 test_ 方法
    failures = []
    total = 0
    for name, obj in list(globals().items()):
        if isinstance(obj, type) and name.startswith("Test"):
            for method_name in dir(obj):
                if method_name.startswith("test_"):
                    method = getattr(obj(), method_name)
                    if callable(method):
                        total += 1
                        try:
                            method()
                            print(f"PASS  {name}.{method_name}")
                        except AssertionError as exc:
                            failures.append((name, method_name, exc))
                            print(f"FAIL  {name}.{method_name}: {exc}")

    print(f"\n{total - len(failures)}/{total} 通过")
    if failures:
        raise SystemExit(1)
